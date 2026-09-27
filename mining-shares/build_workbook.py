"""Build a workbook from a dataset's raw data (no internet needed).

Usage: python build_workbook.py [majors|gold]
  majors: companies.csv + raw/      -> mining_shares_2017_2026.xlsx
  gold:   companies_gold.csv + raw_gold/ + manual_prices_gold.csv
                                    -> gold_miners_2017_2026.xlsx

Prices are shown as they actually traded ("as-traded"). Yahoo serves history
rescaled after splits, bonus issues, stock dividends and demergers; this script
undoes that so a year-end price matches what the stock closed at that day.

The USD tabs, forward tabs and FX forwards are Excel formulas, so editing a
rate on the "Rates 2y" tab updates the whole workbook. The formulas are
evaluated with pycel and their results stored in the file, so viewers that
don't recalculate (phone previews) still show numbers.

Forward price (local): F = S * exp((r - q) * 2)
FX forward (per USD):  X_fwd = X * exp((r - r_usd) * 2)
Forward price (USD):   F / X_fwd   (pence are divided by 100 first)
"""

import math
import re
import shutil
import sys
import zipfile
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).parent
RAW = HERE / "raw"  # set by configure()
FIRST_YEAR, LAST_YEAR = 2017, 2026
YEARS = list(range(FIRST_YEAR, LAST_YEAR + 1))
HORIZON = 2

# Rescaling Yahoo applied that is not listed in its split data (demergers paid
# in shares, rights issues). Each entry: prices before `date` were multiplied
# by `factor` (cumulative, on top of any recorded splits). Factors were
# recovered from the exchange tick size and checked against known closes
# (e.g. Anglo 29-Dec-2023 = 1970.6p, BHP 29-Dec-2017 = A$29.57).
HIDDEN_ADJUSTMENTS = {
    "majors": {
        "BHP.AX": [("2022-05-25", 0.89041)],                          # Woodside in-specie
        "AAL.L": [("2021-06-07", 0.97640), ("2025-06-02", 0.98093)],  # Thungela; Valterra + consolidation
    },
    "etf": {},
    "gold": {
        "SBM.AX": [("2019-05-17", 0.42152), ("2023-07-06", 0.43527)],  # 2019 raise; Genesis in-specie 2023
        "ALK.AX": [("2019-08-14", 0.83129), ("2020-07-20", 0.95870)],  # 2019 raise; ASM demerger 2020
        "RRL.AX": [("2021-04-15", 0.96366)],                           # 2021 entitlement offer
        "RSG.AX": [("2022-11-14", 0.89597)],                           # 2022 entitlement offer
        "VAU.AX": [("2021-03-19", 0.97473)],                           # Red 5 2021 entitlement offer
        "LYC.AX": [("2020-08-19", 0.98635)],                           # 2020 raise
        "GMD.AX": [("2019-08-14", 0.93160), ("2020-06-26", 0.94994), ("2021-11-25", 0.98109)],
        "OBM.AX": [("2020-07-07", 0.90850), ("2022-02-24", 0.91749)],  # pre-2019 (Eastern Goldfields) from manual prices
    },
}
# Cross-check differences that were investigated and are not errors in our data.
EXPLAINED = {
    ("MUX", 2021, "nasdaq"): "Nasdaq labels the 3 Jan 2022 close as 31 Dec 2021; Yahoo 8.87 matches MUX.TO",
    ("OBM.AX", 2017, "cnbc_adjusted"): "CNBC adjusted series includes a later rights issue; CNBC unadjusted matches",
    ("OBM.AX", 2018, "cnbc_adjusted"): "CNBC adjusted series includes a later rights issue; CNBC unadjusted matches",
    ("SBM.AX", 2017, "cnbc_adjusted"): "CNBC adjusted series includes the 2019 capital raise; CNBC unadjusted matches",
    ("SBM.AX", 2018, "cnbc_adjusted"): "CNBC adjusted series includes the 2019 capital raise; CNBC unadjusted matches",
    **{("STLR.TO", y, "tmx"): "One price tick apart (C$0.005) on a sub-C$0.40 stock after TMX restatement"
       for y in (2017, 2018, 2019, 2020)},
}
# Recorded splits that are already part of a HIDDEN_ADJUSTMENTS factor.
IGNORE_SPLITS = {"AAL.L"}
DATASETS = {
    "majors": ("companies.csv", "raw", "mining_shares_2017_2026", None),
    "gold": ("companies_gold.csv", "raw_gold", "gold_miners_2017_2026", "manual_prices_gold.csv"),
    "etf": ("companies_etf.csv", "raw_etf", "pick_etf_2017_2026", None),
}
DATASET = "majors"

# Approximate year-end 2-year government yields in % (Saudi Arabia: SAIBOR-style
# proxy). 2026 repeats 2025 until a real year-end figure exists.
TWO_YEAR_YIELD = {
    "USD": [1.89, 2.49, 1.58, 0.13, 0.73, 4.43, 4.25, 4.24, 3.48],
    "CAD": [1.69, 1.86, 1.70, 0.20, 0.95, 4.05, 3.88, 2.93, 2.60],
    "AUD": [2.00, 1.90, 0.90, 0.10, 0.60, 3.50, 3.70, 3.85, 3.90],
    "GBP": [0.45, 0.75, 0.55, -0.15, 0.70, 3.60, 4.00, 4.40, 3.75],
    "MXN": [7.60, 8.60, 6.90, 4.30, 7.20, 10.50, 10.30, 9.90, 7.30],
    "BRL": [7.50, 7.20, 5.30, 4.30, 11.30, 13.40, 10.30, 15.40, 13.90],
    "SAR": [2.00, 2.90, 2.30, 0.90, 1.20, 5.00, 5.40, 5.20, 4.60],
    "HKD": [1.30, 1.80, 1.70, 0.20, 0.60, 4.00, 3.60, 3.40, 2.60],
    "CNY": [3.50, 2.70, 2.60, 2.70, 2.30, 2.30, 2.20, 1.10, 1.40],
    "IDR": [5.80, 7.70, 5.70, 4.10, 3.90, 5.90, 6.40, 6.90, 5.00],
    "ZAR": [7.00, 7.30, 6.90, 4.50, 5.60, 8.00, 8.20, 7.60, 7.00],
    "SEK": [-0.60, -0.50, -0.30, -0.30, 0.30, 2.70, 2.90, 2.00, 1.90],
    "EUR": [-0.60, -0.60, -0.60, -0.70, -0.60, 2.75, 2.40, 2.10, 2.10],
}
CURRENCIES = list(TWO_YEAR_YIELD)

FONT = Font(name="Arial", size=10)
BOLD = Font(name="Arial", size=10, bold=True)
BLUE = Font(name="Arial", size=10, color="0000FF")
GREEN = Font(name="Arial", size=10, color="008000")
CORRECTED = PatternFill("solid", fgColor="FFF2CC")
MANUAL = PatternFill("solid", fgColor="DDEBF7")
HEADER = PatternFill("solid", fgColor="D9E1F2")


def as_traded_factor(ticker: str, d: pd.DataFrame) -> pd.Series:
    """Multiplier turning Yahoo's rescaled prices into as-traded prices."""
    f = pd.Series(1.0, index=d.index)
    for date, factor in reversed(HIDDEN_ADJUSTMENTS[DATASET].get(ticker, [])):
        f[d.index < pd.Timestamp(date)] = 1 / factor
    if ticker in IGNORE_SPLITS:
        return f
    splits = d["Stock Splits"].fillna(0).replace(0, 1)
    return f * splits[::-1].cumprod()[::-1].shift(-1).fillna(1)


def yearly(prices: pd.DataFrame, ticker: str) -> pd.DataFrame:
    d = prices[prices.ticker == ticker].set_index("date").sort_index()
    d = d[d.index.year >= FIRST_YEAR].dropna(subset=["Close"])
    f = as_traded_factor(ticker, d)
    rows = []
    if d.empty:
        return pd.DataFrame(columns=["last_trading_day", "open", "high", "low", "close", "average_close",
                                     "dividends_in_year", "dividend_yield", "close_yahoo_adjusted"], dtype=object)
    for year, y in d.groupby(d.index.year):
        fy = f.loc[y.index]
        close_adj = y.Close.iloc[-1]
        rows.append({
            "year": year,
            "last_trading_day": y.index[-1].date(),
            "open": y.Open.iloc[0] * fy.iloc[0],
            "high": (y.High * fy).max(),
            "low": (y.Low * fy).min(),
            "close": close_adj * fy.iloc[-1],
            "average_close": (y.Close * fy).mean(),
            "dividends_in_year": (y.Dividends * fy).sum(),
            "dividend_yield": y.Dividends.sum() / close_adj,
            "close_yahoo_adjusted": close_adj,
        })
    return pd.DataFrame(rows).set_index("year")


def fx_year_end() -> pd.DataFrame:
    fx = pd.read_csv(RAW / "fx.csv", parse_dates=["date"]).set_index("date").sort_index()
    ye = fx.ffill().groupby(fx.index.year).last()
    return ye.loc[FIRST_YEAR:LAST_YEAR]


def fx_currency(cur: str) -> str:
    """Excel expression mapping pence (GBp) to the GBP FX row."""
    return 'IF(' + cur + '="GBp","GBP",' + cur + ')'


def style_header(ws, row: int = 1) -> None:
    for c in ws[row]:
        c.font = BOLD
        c.fill = HEADER


def set_widths(ws, widths: dict) -> None:
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def load_manual(path: Path | None) -> pd.DataFrame:
    """Year-end closes researched by hand (delisted stocks, corrections)."""
    if path is None or not path.exists():
        return pd.DataFrame(columns=["ticker", "year", "close", "date", "source"])
    return pd.read_csv(path, dtype={"ticker": str})


def main() -> None:
    global RAW, DATASET, CURRENCIES
    DATASET = sys.argv[1] if len(sys.argv) > 1 else "majors"
    companies_file, raw_dir, out_name, manual_file = DATASETS[DATASET]
    RAW = HERE / raw_dir
    companies = pd.read_csv(HERE / companies_file)
    listed = (companies.dropna(subset=["ticker"]).drop_duplicates("ticker")
              .sort_values("ticker").reset_index(drop=True))
    prices = pd.read_csv(RAW / "prices.csv", parse_dates=["date"])
    targets = pd.read_csv(RAW / "targets.csv").set_index("ticker")
    fx_ye = fx_year_end()
    CURRENCIES = ["USD"] + [c for c in TWO_YEAR_YIELD if c != "USD" and c in fx_ye.columns]
    data = {t: yearly(prices, t) for t in listed.ticker}
    manual = load_manual(HERE / manual_file if manual_file else None)
    manual_cells = {}
    for m in manual.itertuples(index=False):
        d = data[m.ticker]
        d.loc[int(m.year), "close"] = float(m.close)
        d.loc[int(m.year), "last_trading_day"] = m.date
        manual_cells[(m.ticker, int(m.year))] = f"Researched close ({m.date}). Source: {m.source}"
        data[m.ticker] = d.sort_index()

    wb = Workbook()
    n = len(listed)
    first, last = 2, n + 1  # data rows on every per-ticker sheet
    year_col = {y: get_column_letter(4 + i) for i, y in enumerate(YEARS)}
    ncols = 3 + len(YEARS)

    def ticker_sheet(title: str, header_years: list) -> "Worksheet":
        ws = wb.create_sheet(title)
        ws.append(["ticker", "company", "currency"] + [str(y) for y in header_years])
        style_header(ws)
        for i, c in listed.iterrows():
            r = first + i
            ws.cell(r, 1, c.ticker)
            ws.cell(r, 2, c.listed_parent)
            ws.cell(r, 3, c.currency)
        set_widths(ws, {"A": 13, "B": 32, "C": 9, **{get_column_letter(4 + j): 11 for j in range(len(header_years))}})
        ws.freeze_panes = "D2"
        return ws

    # Entities
    ws = wb.active
    ws.title = "Entities"
    ws.append(list(companies.columns))
    style_header(ws)
    for row in companies.itertuples(index=False):
        ws.append([None if pd.isna(v) else v for v in row])
    set_widths(ws, {"A": 40, "B": 34, "C": 13, "D": 14, "E": 9, "F": 60})

    # Rates 2y (inputs)
    rates = wb.create_sheet("Rates 2y")
    rates.append(["currency"] + [str(y) for y in YEARS])
    style_header(rates)
    for cur in CURRENCIES:
        vals = TWO_YEAR_YIELD[cur]
        rates.append([cur] + [vals[min(i, len(vals) - 1)] / 100 for i in range(len(YEARS))])
    for row in rates.iter_rows(min_row=2, min_col=2):
        for c in row:
            c.font, c.number_format = BLUE, "0.00%"
    rates.cell(len(CURRENCIES) + 3, 1, "Approximate year-end 2-year government bond yields (SAR: interbank "
               "proxy). Entered from memory, not from a data source: replace with exact figures if needed. "
               "2026 repeats 2025.").font = FONT
    set_widths(rates, {"A": 10, **{year_col[y]: 9 for y in YEARS}})
    rates_rng = f"'Rates 2y'!$A$2:$A${len(CURRENCIES) + 1}"

    def rate_ref(cur_expr: str, year: int) -> str:
        col = get_column_letter(2 + YEARS.index(year))
        return f"INDEX('Rates 2y'!${col}$2:${col}${len(CURRENCIES) + 1},MATCH({cur_expr},{rates_rng},0))"

    # FX year-end (data)
    fxs = wb.create_sheet("FX year-end")
    fxs.append(["currency"] + [str(y) for y in YEARS])
    style_header(fxs)
    fxs.append(["USD"] + [1] * len(YEARS))
    for cur in CURRENCIES[1:]:
        fxs.append([cur] + [round(float(fx_ye.loc[y, cur]), 6) for y in YEARS])
    fxs.cell(len(CURRENCIES) + 3, 1, "Units of currency per 1 USD at the last trading day of each year "
             "(2026: latest). Source: Yahoo Finance (CAD=X, GBP=X, ...).").font = FONT
    set_widths(fxs, {"A": 10, **{year_col[y]: 12 for y in YEARS}})
    fx_rng = f"'FX year-end'!$A$2:$A${len(CURRENCIES) + 1}"

    def fx_ref(sheet: str, cur_expr: str, year: int) -> str:
        col = get_column_letter(2 + YEARS.index(year))
        return f"INDEX('{sheet}'!${col}$2:${col}${len(CURRENCIES) + 1},MATCH({cur_expr},'{sheet}'!$A$2:$A${len(CURRENCIES) + 1},0))"

    # FX 2y forward (formulas), columns = year the forward was set
    fxf = wb.create_sheet("FX 2y forward")
    fxf.append(["currency"] + [f"{y}->{y + HORIZON}" for y in YEARS])
    style_header(fxf)
    for i, cur in enumerate(CURRENCIES):
        r = 2 + i
        fxf.cell(r, 1, cur)
        for j, y in enumerate(YEARS):
            spot = f"'FX year-end'!{get_column_letter(2 + j)}{r}"
            r_ccy, r_usd = rate_ref(f"$A{r}", y), rate_ref('"USD"', y)
            fxf.cell(r, 2 + j, f"={spot}*EXP(({r_ccy}-{r_usd})*{HORIZON})")
    fxf.cell(len(CURRENCIES) + 3, 1, "Covered-interest-parity forward: spot x exp((r_ccy - r_usd) x 2). "
             "Column '2017->2019' = rate agreed at end-2017 for delivery at end-2019.").font = FONT
    set_widths(fxf, {"A": 10, **{year_col[y]: 12 for y in YEARS}})

    # Year-end close (local, data)
    close = ticker_sheet("Year-end close", YEARS)
    close.cell(1, 3).value = "currency"
    for i, c in listed.iterrows():
        d = data[c.ticker]
        for y in YEARS:
            if y in d.index:
                cell = close[f"{year_col[y]}{first + i}"]
                cell.value = round(float(d.loc[y, "close"]), 4)
                adj = float(d.loc[y, "close_yahoo_adjusted"])
                if (c.ticker, y) in manual_cells:
                    cell.fill = MANUAL
                    cell.comment = Comment(manual_cells[(c.ticker, y)], "Claude")
                elif abs(cell.value / adj - 1) > 0.001:
                    cell.fill = CORRECTED
                    cell.comment = Comment(f"As traded. Yahoo shows {adj:.4f} (history rescaled after a "
                                           "split / bonus issue / demerger).", "Claude")

        raw = prices[(prices.ticker == c.ticker) & (prices["Stock Splits"].fillna(0) != 0)]
        for date, ratio in zip(raw.date, raw["Stock Splits"]):
            if c.ticker in IGNORE_SPLITS or date.year < FIRST_YEAR:
                continue
            n_for_1 = round(1 / ratio if ratio < 1 else ratio, 2)
            kind = "share consolidation" if ratio < 1 else "split"
            y = date.year
            cell = close[f"{year_col[y]}{first + i}"] if y in year_col else None
            if cell is not None and cell.value is not None:
                note = (f"{n_for_1:g}-for-1 {kind} on {date.date()}: earlier years are pre-{kind.split()[-1]} "
                        "prices, so a jump here is not a real price move.")
                cell.comment = Comment(note if cell.comment is None else cell.comment.text + "\n" + note, "Claude")
                cell.font = Font(name="Arial", size=10, bold=True)
        if d.empty:
            close.cell(first + i, 1).comment = Comment("No free price history found for this delisted stock.",
                                                       "Claude")

    # Dividend yield (data)
    dy = ticker_sheet("Dividend yield", YEARS)
    for i, c in listed.iterrows():
        d = data[c.ticker]
        for y in YEARS:
            if y in d.index and not pd.isna(d.loc[y, "dividend_yield"]):
                cell = dy[f"{year_col[y]}{first + i}"]
                cell.value = round(float(d.loc[y, "dividend_yield"]), 6)
                cell.number_format = "0.00%"

    # Year-end close USD (formulas)
    close_usd = ticker_sheet("Year-end close USD", YEARS)
    for i in range(n):
        r = first + i
        close_usd.cell(r, 3, "USD")
        for y in YEARS:
            col = year_col[y]
            local = f"'Year-end close'!{col}{r}"
            cur = f"'Year-end close'!$C{r}"
            fx = fx_ref("FX year-end", fx_currency(cur), y)
            close_usd[f"{col}{r}"] = f'=IF({local}="","",{local}/IF({cur}="GBp",100,1)/{fx})'

    # 2y forward (local) - columns are the target year
    target_years = [y + HORIZON for y in YEARS]
    fwd = ticker_sheet("2y forward", target_years)
    for i in range(n):
        r = first + i
        for y in YEARS:
            col, tcol = year_col[y], year_col[y]  # same position: column j = set at YEARS[j], for YEARS[j]+2
            s = f"'Year-end close'!{col}{r}"
            q = f"'Dividend yield'!{col}{r}"
            cur = fx_currency(f"$C{r}")
            fwd[f"{tcol}{r}"] = f'=IF({s}="","",{s}*EXP(({rate_ref(cur, y)}-{q})*{HORIZON}))'

    # 2y forward USD
    fwd_usd = ticker_sheet("2y forward USD", target_years)
    for i in range(n):
        r = first + i
        fwd_usd.cell(r, 3, "USD")
        for y in YEARS:
            col = year_col[y]
            f_local = f"'2y forward'!{col}{r}"
            cur = f"'2y forward'!$C{r}"
            fx = fx_ref("FX 2y forward", fx_currency(cur), y)
            fwd_usd[f"{col}{r}"] = f'=IF({f_local}="","",{f_local}/IF({cur}="GBp",100,1)/{fx})'

    for ws in (close, dy, close_usd, fwd, fwd_usd):
        for row in ws.iter_rows(min_row=first, max_row=last, min_col=4, max_col=ncols):
            for c in row:
                if ws is not dy:
                    c.number_format = "#,##0.00"
        for row in ws.iter_rows(min_row=first, max_row=last, max_col=3):
            for c in row:
                c.font = FONT
    for ws in (close_usd, fwd, fwd_usd):
        for row in ws.iter_rows(min_row=first, max_row=last, min_col=4, max_col=ncols):
            for c in row:
                c.font = GREEN if ws is close_usd else FONT

    notes = {
        close: "Last close of each year, in the currency the stock trades in (GBp = pence). 2026 = latest close. "
               "Highlighted cells were corrected to the as-traded price; hover for Yahoo's rescaled figure.",
        close_usd: "Year-end close / year-end FX rate (pence divided by 100 first).",
        fwd: "Price a 2-year future set at the end of (column year - 2) would have had: "
             "close x exp((2y rate - dividend yield) x 2). Carry only, not a forecast.",
        fwd_usd: "2y forward (local) / 2-year FX forward from the 'FX 2y forward' tab.",
        dy: "Dividends paid during the year / year-end close.",
    }
    for ws, text in notes.items():
        ws.cell(last + 2, 1, text).font = FONT

    # Detail (data)
    det = wb.create_sheet("Detail")
    cols = ["ticker", "company", "currency", "year", "last_trading_day", "open", "high", "low", "close",
            "average_close", "dividends_in_year", "dividend_yield", "rate_2y", "forward_2y_price",
            "forward_for_year", "actual_close_in_forward_year", "forward_error"]
    det.append(cols)
    style_header(det)
    csv_rows = []
    for _, c in listed.iterrows():
        d = data[c.ticker]
        for y in d.index:
            cur = c.currency.replace("GBp", "GBP")
            r = TWO_YEAR_YIELD[cur][min(y - FIRST_YEAR, 8)] / 100
            row = d.loc[y]
            q = 0.0 if pd.isna(row.dividend_yield) else row.dividend_yield
            f = row.close * math.exp((r - q) * HORIZON)
            later = d.loc[y + HORIZON] if y + HORIZON in d.index else None
            actual = later.close if later is not None else None
            # Error on Yahoo's consistently rescaled series, so a split in between doesn't distort it.
            err = None
            if later is not None and not pd.isna(row.close_yahoo_adjusted) and not pd.isna(later.close_yahoo_adjusted):
                f_adj = row.close_yahoo_adjusted * math.exp((r - q) * HORIZON)
                err = later.close_yahoo_adjusted / f_adj - 1
            elif later is not None:
                err = later.close / f - 1
            vals = [c.ticker, c.listed_parent, c.currency, y, row.last_trading_day, row.open, row.high,
                    row.low, row.close, row.average_close, row.dividends_in_year, row.dividend_yield, r, f,
                    y + HORIZON, actual, err]
            vals = [None if isinstance(v, float) and math.isnan(v) else round(v, 4) if isinstance(v, float) else v
                    for v in vals]
            det.append(vals)
            csv_rows.append(dict(zip(cols, vals)))
    for row in det.iter_rows(min_row=2):
        for c in row:
            c.font = FONT
        for idx in (11, 12, 16):
            row[idx].number_format = "0.0%"
    det.freeze_panes = "E2"
    set_widths(det, {"A": 13, "B": 32, "E": 14})
    pd.DataFrame(csv_rows).to_csv(HERE / f"{out_name}.csv", index=False)

    # Analyst targets (now)
    at = wb.create_sheet("Analyst targets (now)")
    at.append(["ticker", "company", "currency", "current", "mean", "median", "low", "high"])
    style_header(at)
    for _, c in listed.iterrows():
        t = targets.loc[c.ticker] if c.ticker in targets.index else None
        at.append([c.ticker, c.listed_parent, c.currency] +
                  [None if t is None or pd.isna(t[k]) else round(float(t[k]), 4)
                   for k in ("current", "mean", "median", "low", "high")])
    at.cell(n + 3, 1, "12-month consensus price targets as of the data download date. Source: Yahoo Finance.").font = FONT
    set_widths(at, {"A": 13, "B": 32})

    order = ["Entities", "Year-end close", "Year-end close USD", "2y forward", "2y forward USD",
             "Dividend yield", "Rates 2y", "FX year-end", "FX 2y forward", "Detail", "Analyst targets (now)"]
    if verification_sheet(wb, data, prices):
        order.append("Verification")
    wb._sheets = [wb[s] for s in order]
    for ws in wb:
        for row in ws.iter_rows():
            for c in row:
                if c.font == Font():
                    c.font = FONT
    out = HERE / f"{out_name}.xlsx"
    wb.save(out)
    store_cached_values(out, [ws.title for ws in wb])
    print(f"Wrote {out.name}")


def verification_sheet(wb: Workbook, data: dict, prices: pd.DataFrame) -> bool:
    """Compare our year-end closes with sources independent of Yahoo (raw/verify.csv)."""
    path = RAW / "verify.csv"
    if not path.exists():
        return False
    v = pd.read_csv(path)
    v = v[(v.year >= FIRST_YEAR) & v.ticker.isin(data)]
    # CNBC's BGL-AU is a different security from Bellevue Gold (ASX: BGL).
    v = v[~((v.ticker == "BGL.AX") & v.source.str.startswith("cnbc"))]
    p = prices[prices.date.dt.year >= FIRST_YEAR].sort_values("date")
    restated = p.groupby(["ticker", p.date.dt.year]).Close.last()
    rows = []
    for r in v.itertuples():
        d = data[r.ticker]
        ours = d.close.get(r.year) if r.year in d.index else None
        if ours is None or pd.isna(ours):
            continue
        yahoo = restated.get((r.ticker, r.year))

        def same(a, b):
            return b is not None and not pd.isna(b) and (abs(a / b - 1) <= 0.01 or abs(a - b) <= 0.0051)
        if (r.ticker, r.year, r.source) in EXPLAINED:
            status = "explained: " + EXPLAINED[(r.ticker, r.year, r.source)]
        elif same(r.close, ours):
            status = "match"
        elif same(r.close, yahoo):
            status = "match (source restated for a later split/consolidation)"
        else:
            status = "MISMATCH"
        rows.append([r.ticker, r.year, r.source, r.date, r.close, round(float(ours), 4), status])
    ws = wb.create_sheet("Verification")
    df = pd.DataFrame(rows, columns=["ticker", "year", "source", "source_date", "source_close", "our_close",
                                     "status"])
    counts = df.status.str.split(":").str[0].value_counts()
    checked = df.ticker.nunique()
    ws.append(["Cross-check of year-end closes against sources independent of Yahoo Finance"])
    ws.append([f"{len(df)} comparisons across {checked} tickers: "
               + ", ".join(f"{k}: {v}" for k, v in counts.items())])
    unchecked = sorted(t for t, d in data.items() if not d.empty and t not in set(df.ticker))
    ws.append(["Not independently checked (no second source reachable): " + (", ".join(unchecked) or "none")])
    ws.append([])
    ws.append(list(df.columns))
    style_header(ws, 5)
    for row in df.sort_values(["status", "ticker", "year"]).itertuples(index=False):
        ws.append(list(row))
        if row.status == "MISMATCH":
            for c in ws[ws.max_row]:
                c.fill = CORRECTED
    ws.cell(1, 1).font = BOLD
    ws.freeze_panes = "A6"
    set_widths(ws, {"A": 13, "C": 16, "D": 12, "G": 52})
    return True


def store_cached_values(path: Path, sheet_titles: list) -> None:
    """Evaluate every formula with pycel and write the results into the file."""
    from pycel import ExcelCompiler

    xl = ExcelCompiler(filename=str(path))
    tmp = path.with_suffix(".tmp")
    with zipfile.ZipFile(path) as src, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
        for item in src.infolist():
            data = src.read(item.filename)
            m = re.fullmatch(r"xl/worksheets/sheet(\d+)\.xml", item.filename)
            if m:
                title = sheet_titles[int(m.group(1)) - 1]

                def fill(cell: re.Match) -> str:
                    value = xl.evaluate(f"'{title}'!{cell.group(2)}")
                    if isinstance(value, str) and value.startswith("#"):
                        raise ValueError(f"{title}!{cell.group(2)} evaluates to {value}")
                    if value in ("", None):
                        return f'<c r="{cell.group(2)}"{cell.group(3)} t="str"><f>{cell.group(4)}</f><v></v></c>'
                    return f'<c r="{cell.group(2)}"{cell.group(3)}><f>{cell.group(4)}</f><v>{float(value)!r}</v></c>'

                text = re.sub(r'(<c r="([A-Z]+\d+)"((?: s="\d+")?)><f>(.*?)</f><v></v></c>)', fill, data.decode())
                data = text.encode()
            dst.writestr(item, data)
    shutil.move(tmp, path)


if __name__ == "__main__":
    main()
